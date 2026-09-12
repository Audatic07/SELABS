from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Architecture_Justification.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="8"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_font(run, name="Aptos", size=10.5, bold=False, color="1E293B"):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_labeled_paragraph(doc, label, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(label + " ")
    set_font(r, bold=True, color="173B66")
    r = p.add_run(body)
    set_font(r)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.62)
section.bottom_margin = Inches(0.58)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(30, 41, 59)

title_style = doc.styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title_style.font.size = Pt(23)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style_ppr = title_style._element.get_or_add_pPr()
title_style_border = title_style_ppr.find(qn("w:pBdr"))
if title_style_border is not None:
    title_style_ppr.remove(title_style_border)

heading = doc.styles["Heading 1"]
heading.font.name = "Aptos Display"
heading._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
heading._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
heading.font.size = Pt(12.5)
heading.font.bold = True
heading.font.color.rgb = RGBColor(0, 0, 0)
heading.paragraph_format.space_before = Pt(7)
heading.paragraph_format.space_after = Pt(3)
heading.paragraph_format.keep_with_next = True

title = doc.add_paragraph(style="Title")
title.paragraph_format.space_after = Pt(2)
title.add_run("Coffee Kiosk Architecture Justification")
title_ppr = title._p.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(10)
r = sub.add_run("Lab 3  |  Component Modelling and Architectural Pattern Selection")
set_font(r, size=9.5, bold=True, color="607089")

doc.add_heading("Architecture Selection", level=1)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(6)
p.paragraph_format.line_spacing = 1.08
r = p.add_run("We chose Layered Architecture for the Self-Service Coffee Kiosk System. ")
set_font(r, bold=True, color="173B66")
r = p.add_run(
    "The kiosk is a compact application with one user interface, a small ordering workflow, "
    "credit-card payment, local menu data, and receipt-printer integration. Separating these "
    "responsibilities into presentation, business, and data/device layers provides enough "
    "structure without the deployment and networking overhead of distributed services."
)
set_font(r)

doc.add_heading("Architectural Style Analysis", level=1)
table = doc.add_table(rows=1, cols=3)
table.autofit = False
table.columns[0].width = Inches(1.15)
table.columns[1].width = Inches(3.45)
table.columns[2].width = Inches(2.15)
headers = ["Style", "Fit for this scenario", "Main tradeoff"]
for i, text in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.width = table.columns[i].width
    set_cell_shading(cell, "173B66")
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i else WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, size=9.5, bold=True, color="FFFFFF")

rows = [
    ("Layered", "Strong fit: simple module boundaries and direct local calls suit one kiosk.", "Some cross-layer call overhead."),
    ("Microservices", "Supports independent scaling, but the kiosk has too few functions to justify distribution.", "Network, deployment, and consistency complexity."),
    ("Client-Server", "Central control could help a fleet, but the stated scenario is a self-contained kiosk.", "Server availability becomes a bottleneck and failure point."),
]
for row_index, values in enumerate(rows, 1):
    cells = table.add_row().cells
    for i, text in enumerate(values):
        cells[i].width = table.columns[i].width
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cells[i], top=85, bottom=85)
        if row_index % 2 == 0:
            set_cell_shading(cells[i], "F2F6FA")
        p = cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        set_font(r, size=9.1, bold=(i == 0), color="1E293B")
set_table_borders(table)

doc.add_heading("Two Scenario Specific Reasons", level=1)
add_labeled_paragraph(
    doc,
    "1  Clear responsibility boundaries.",
    "The Touchscreen UI can change independently of the Order Manager, while menu storage, payment authorization, and printer control remain behind stable interfaces. This makes the small system easier to test, replace, and maintain.",
)
add_labeled_paragraph(
    doc,
    "2  Practical hardware and data integration.",
    "The business layer coordinates an order without knowing SQL details or printer commands. The Menu Repository and Receipt Printer Adapter isolate those technologies, reducing the impact of a database or printer-model change.",
)

doc.add_heading("Security Advantage", level=1)
add_labeled_paragraph(
    doc,
    "Payment isolation.",
    "Only the Order Manager can request payment through the Payment API. The Touchscreen UI and Menu Repository never receive raw card data; the Payment Service validates the request and returns only an authorization result and transaction reference. This narrows the sensitive-data boundary and supports input validation and least-privilege access.",
)

doc.add_heading("Performance Benefit", level=1)
add_labeled_paragraph(
    doc,
    "Low-latency local flow.",
    "Presentation, ordering, menu lookup, and printer coordination can use in-process calls or local adapters. Avoiding multiple network hops gives responsive touch interactions and predictable checkout time. Frequently read menu and pricing data can also be cached in the business layer while the repository remains the source of truth.",
)

doc.add_heading("Component and Interface Summary", level=1)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
p.paragraph_format.line_spacing = 1.05
r = p.add_run(
    "Five components are used: Touchscreen UI, Order Manager, Payment Service, Menu Repository, "
    "and Receipt Printer Adapter. Their four assembly interfaces are Order API, Payment API, "
    "Menu Query Interface, and Printer Port."
)
set_font(r, size=9.6)

doc.core_properties.title = "Coffee Kiosk Architecture Justification"
doc.core_properties.subject = "Lab 3 Component Modelling and Architectural Pattern Selection"
doc.core_properties.author = "Audatic07"
doc.core_properties.keywords = "UML, component diagram, layered architecture, coffee kiosk"

doc.save(OUTPUT)
print(OUTPUT)
